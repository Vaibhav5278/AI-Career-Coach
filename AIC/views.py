from users1.models import *
import os

def index(request):
    """
    Render the home page with FAQ section.
    """

    # You can later fetch these from the database if needed.
    faqs = [
        {
            "question": "What is AI Career Coach?",
            "answer": "AI Career Coach is a platform that helps you plan your career using AI-based guidance, resources, and expert advice."
        },
        {
            "question": "Is the service free?",
            "answer": "Yes, the basic AI coaching features are free. Advanced mentorship options may be added later."
        },
        {
            "question": "How does it suggest career paths?",
            "answer": "The system analyzes your skills, interests, and job market trends using machine learning models."
        },
        {
            "question": "Can I get personalized guidance?",
            "answer": "Absolutely. Create a profile and the AI engine will tailor recommendations to your background and goals."
        },
        {
            "question": "What technologies are used?",
            "answer": "The platform uses Python, Django, and AI/ML models for natural language understanding and career predictions."
        }
    ]

    return render(request, "index.html", {"faqs": faqs})

def login(request):
    return render(request, "login.html")


from django.contrib import messages

from django.db import IntegrityError

def register(request):
    if request.method == "POST":
        username = request.POST.get("username")
        email = request.POST.get("email")
        password = request.POST.get("password")
        confirm_password = request.POST.get("confirm_password")

        if password != confirm_password:
            messages.error(request, "Passwords do not match")
            return redirect("register")

        try:
            user = User.objects.create(
                username=username,
                email=email,
                password=password,
                gender=None,
                role="User"
            )
            messages.success(request, "Account created successfully!")
            return redirect("login")

        except IntegrityError:
            messages.error(request, "Email already exists")
            return redirect("register")

    return render(request, "register.html")


def dashboard(request):
    return render(request, "dashboard.html")

def entry_cv(request):
    return render(request, "entry_cv.html")


from django.http import HttpResponse
from django.template.loader import render_to_string
from xhtml2pdf import pisa
from docx import Document


# --------------------------
# Entry Resume Form
# --------------------------

from django.shortcuts import redirect

from django.shortcuts import render
from django.contrib.auth.decorators import login_required

#@login_required
from django.shortcuts import render, redirect
from django.views.decorators.csrf import csrf_exempt
from django.http import HttpResponse
import datetime
def entry_resume(request):
    """
    Handles resume form submission and saves data in session
    """

    if request.method == "POST":

        # --- Basic Info ---
        name = request.POST.get("name")
        email = request.POST.get("email")
        phone = request.POST.get("phone")
        linkedin = request.POST.get("linkedin")
        twitter = request.POST.get("twitter")
        summary = request.POST.get("summary")
        skills = request.POST.get("skills")
        template = request.POST.get("template")

        # --- Education ---
        edu_titles = request.POST.getlist("edu_title")
        edu_schools = request.POST.getlist("edu_school")
        edu_starts = request.POST.getlist("edu_start")
        edu_ends = request.POST.getlist("edu_end")
        edu_descs = request.POST.getlist("edu_desc")

        educations = []

        for i in range(len(edu_titles)):
            if edu_titles[i].strip():
                educations.append({
                    "title": edu_titles[i],
                    "school": edu_schools[i],
                    "start": edu_starts[i],
                    "end": edu_ends[i],
                    "desc": edu_descs[i]
                })

        # --- Projects ---
        proj_titles = request.POST.getlist("project_title[]")
        proj_techs = request.POST.getlist("project_tech[]")
        proj_descs = request.POST.getlist("project_desc[]")

        projects = []

        for i in range(len(proj_titles)):
            if proj_titles[i].strip():
                projects.append({
                    "title": proj_titles[i],
                    "tech": proj_techs[i],
                    "desc": proj_descs[i]
                })

        # --- Work Experience ---
        work_titles = request.POST.getlist("work_title[]")
        work_companies = request.POST.getlist("work_company[]")
        work_starts = request.POST.getlist("work_start[]")
        work_ends = request.POST.getlist("work_end[]")
        work_descs = request.POST.getlist("work_desc[]")

        work_experiences = []

        for i in range(len(work_titles)):
            if work_titles[i].strip():
                work_experiences.append({
                    "title": work_titles[i],
                    "company": work_companies[i],
                    "start": work_starts[i],
                    "end": work_ends[i],
                    "desc": work_descs[i]
                })

        # --- Certifications ---
        cert_names = request.POST.getlist("cert_name[]")
        cert_orgs = request.POST.getlist("cert_org[]")
        cert_dates = request.POST.getlist("cert_date[]")

        certifications = []

        for i in range(len(cert_names)):
            cert_name = cert_names[i]   # ✅ FIXED (was: name = cert_names[i])
            org = cert_orgs[i] if i < len(cert_orgs) else ""
            date = cert_dates[i] if i < len(cert_dates) else ""

            if cert_name.strip():
                certifications.append({
                    "name": cert_name,
                    "org": org,
                    "date": date,
                })

        # --- Extra Activities ---
        extra_activities = request.POST.getlist("extra[]")

        activities = [act for act in extra_activities if act.strip()]

        # --- Prepare Data Object ---
        data = {
            "name": name,
            "email": email,
            "phone": phone,
            "linkedin": linkedin,
            "twitter": twitter,
            "summary": summary,
            "skills": skills.split(",") if skills else [],
            "education": educations,
            "projects": projects,
            "work": work_experiences,
            "certifications": certifications,
            "extra_activities": activities
        }

        # --- Save to Session ---
        request.session["resume_data"] = data
        request.session["template_choice"] = template

        # --- Redirect to Preview ---
        return redirect("resume_preview")

    return render(request, "entry_resume.html")
# --------------------------
# Resume Preview Page
# --------------------------
def resume_preview(request):

    data = request.session.get("resume_data")
    template_name = request.session.get("template_choice", "template1")

    if not data:
        return redirect("entry_resume")

    template_path = f"resume_templates/{template_name}.html"

    return render(request, template_path, {"data": data})

# --------------------------
# Template Selection Page
# --------------------------
def choose_template(request):
    if "resume_data" not in request.session:
        return redirect("entry_resume")

    if request.method == "POST":
        request.session["template_choice"] = request.POST.get("template", "template1")
        return redirect("resume_preview")

    return render(request, "choose_template.html")


# --------------------------
# PDF Download
# --------------------------
def download_resume_pdf(request):
    data = request.session.get("resume_data")
    template_name = request.session.get("template_choice", "template1")

    if not data:
        return redirect("entry_resume")

    html = render_to_string(f"resume_templates/{template_name}.html", {"data": data})

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="Resume.pdf"'

    pisa_status = pisa.CreatePDF(html, dest=response)

    if pisa_status.err:
        return HttpResponse("Error generating PDF")

    return response


# --------------------------
# DOCX Download
# --------------------------
def download_resume_docx(request):
    data = request.session.get("resume_data")

    if not data:
        return redirect("entry_resume")

    doc = Document()

    doc.add_heading(data.get("name", ""), 0)
    doc.add_paragraph(f"Email: {data.get('email', '')}")
    doc.add_paragraph(f"Phone: {data.get('phone', '')}")
    doc.add_paragraph(f"LinkedIn: {data.get('linkedin', '')}")
    doc.add_paragraph(f"Twitter: {data.get('twitter', '')}")

    doc.add_heading("Summary", 1)
    doc.add_paragraph(data.get("summary", ""))

    doc.add_heading("Skills", 1)
    doc.add_paragraph(data.get("skills", ""))

    doc.add_heading("Education", 1)
    doc.add_paragraph(
        f"{data.get('edu_title_1', '')} at {data.get('edu_school_1', '')} "
        f"({data.get('edu_start_1', '')} - {data.get('edu_end_1', '')})"
    )

    doc.add_heading("Experience", 1)
    doc.add_paragraph(f"{data.get('work_title_1', '')} at {data.get('work_company_1', '')}")

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    response["Content-Disposition"] = 'attachment; filename=Resume.docx'

    doc.save(response)
    return response



# views.py


# views.py

from django.shortcuts import render

from django.shortcuts import render

def cover_preview(request):
    if request.method == 'POST':
        data = request.POST
        template_choice = data.get('template')

        context = {
            'title': data.get('title'),
            'name': data.get('name'),
            'address': data.get('address'),
            'email': data.get('email'),
            'phone': data.get('phone'),
            'linkedin': data.get('linkedin'),
            'date': data.get('date'),
            'company': data.get('company'),
            'company_address': data.get('company_address'),
            'position': data.get('position'),
            'hiring_manager': data.get('hiring_manager'),
            'intro': data.get('intro'),
            'body': data.get('body'),
            'closing': data.get('closing'),
            'signature': data.get('signature'),
        }

        # Choose template dynamically
        if template_choice == 'template1':
            return render(request, 'cover_letters/cv_template1.html', context)
        elif template_choice == 'template2':
            return render(request, 'cover_letters/template2.html', context)
        elif template_choice == 'template3':
            return render(request, 'cover_letters/template3.html', context)
        elif template_choice == 'template4':
            return render(request, 'cover_letters/template4.html', context)
        else:
            return render(request, 'cover_letters/cv_template1.html', context)

    # --- Handle GET request: show empty/default form ---
    return render(request, 'cover_letters/cover_form.html')


from django.template.loader import render_to_string
from xhtml2pdf import pisa
from django.http import HttpResponse, HttpResponseRedirect

def download_cover_pdf(request):
    data = request.session.get("cover_data")
    template_name = request.session.get("cover_template_choice", "template1")

    if not data:
        return redirect("cover_preview")  # Or your cover letter entry page

    html = render_to_string(f"cover_letters/{template_name}.html", data)

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="Cover_Letter.pdf"'

    pisa_status = pisa.CreatePDF(html, dest=response)

    if pisa_status.err:
        return HttpResponse("Error generating PDF")

    return response


from docx import Document

def download_cover_docx(request):
    data = request.session.get("cover_data")

    if not data:
        return redirect("cover_preview")

    doc = Document()
    doc.add_heading(data.get("title", "Cover Letter"), 0)
    doc.add_paragraph(f"Date: {data.get('date', '')}")
    doc.add_paragraph(f"To: {data.get('hiring_manager', '')}")
    doc.add_paragraph(f"Company: {data.get('company', '')}")
    doc.add_paragraph(f"Company Address: {data.get('company_address', '')}")
    doc.add_paragraph(f"Position: {data.get('position', '')}")

    doc.add_paragraph("\n" + data.get("intro", ""))
    doc.add_paragraph("\n" + data.get("body", ""))
    doc.add_paragraph("\n" + data.get("closing", ""))
    doc.add_paragraph("\n" + data.get("signature", ""))

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    response["Content-Disposition"] = 'attachment; filename=Cover_Letter.docx'

    doc.save(response)
    return response


def cv_template1(request):
    return render(request, 'cover_letters/cv_template1.html')

def learning_path(request):
    return render(request,'learning_path.html')


# from openai import OpenAI
# import os
#
# # Initialize OpenAI client
# client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))  # Store your key in .env file
# from openai import OpenAI
# import os
# from dotenv import load_dotenv
#
# load_dotenv()
# client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
#
# from django.http import JsonResponse
# from django.views.decorators.csrf import csrf_exempt
# from openai import OpenAI
# import os, json
# from dotenv import load_dotenv
#
# load_dotenv()
# client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
#
# @csrf_exempt
# def generate_summary(request):
#     if request.method == "POST":
#         try:
#             data = json.loads(request.body)
#             keyword = data.get("keyword", "").strip()
#             if not keyword:
#                 return JsonResponse({"error": "No keyword provided"}, status=400)
#
#             prompt = f"Write a short professional resume description for: {keyword}"
#
#             response = client.chat.completions.create(
#                 model="gpt-4o-mini",
#                 messages=[
#                     {"role": "system", "content": "You are a helpful assistant that writes resume content."},
#                     {"role": "user", "content": prompt},
#                 ],
#                 max_tokens=150
#             )
#
#             description = response.choices[0].message.content.strip()
#             return JsonResponse({"description": description})
#
#         except Exception as e:
#             return JsonResponse({"error": str(e)}, status=500)
#
#     return JsonResponse({"error": "Invalid request"}, status=405)

def about(request):
    return render(request, 'about.html')

def contact(request):
    return render(request, 'contact.html')


import json
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_protect
from huggingface_hub import InferenceClient

# Initialize the AI Client
# Replace with your actual Hugging Face token

import os
from huggingface_hub import InferenceClient

client = InferenceClient(api_key=os.getenv("HF_TOKEN"))



def generate_summary(request):
    # --- 📝 LOGGING: See exactly what is arriving ---
    print("\n--- [DEBUG] Incoming Request ---")
    print(f"Content-Type: {request.content_type}")
    print(f"Raw Body: {request.body.decode('utf-8')}")

    if request.method != "POST":
        return JsonResponse({"error": "Only POST allowed"}, status=405)

    try:
        # ✅ STEP 1: Parse the data safely
        if "application/json" in request.content_type:
            data = json.loads(request.body.decode("utf-8"))
        else:
            data = request.POST

        # ✅ STEP 2: Get the text from the 'summary' key
        user_text = data.get("summary", "").strip()
        print(f"Extracted Text: {user_text}")

        if not user_text:
            return JsonResponse({"error": "The summary box is empty!"}, status=400)

        # ✅ STEP 3: Call Hugging Face
        prompt = f"Improve this professional resume summary: '{user_text}'. Return ONLY the improved version."

        completion = client.chat.completions.create(
            model="meta-llama/Meta-Llama-3-8B-Instruct",
            messages=[
                {"role": "system", "content": "You are a professional resume writer."},
                {"role": "user", "content": prompt}
            ],
            max_tokens=250,
            temperature=0.7
        )

        ai_summary = completion.choices[0].message.content.strip()
        print(f"AI Result: {ai_summary}")

        return JsonResponse({"summary": ai_summary})

    except Exception as e:
        # This will print the exact line and reason for the 500 error in your terminal
        import traceback
        print("\n--- 🔥 SERVER ERROR ---")
        print(traceback.format_exc())
        return JsonResponse({"error": str(e)}, status=500)



# education
def generate_education(request):

    if request.method != "POST":
        return JsonResponse({"error": "Only POST allowed"}, status=405)

    try:
        data = json.loads(request.body.decode("utf-8"))
        user_text = data.get("content", "").strip()

        if not user_text:
            return JsonResponse({"error": "Education is empty"}, status=400)

        prompt = f"""
        Improve this education description professionally:

        {user_text}

        Return ONLY improved version.
        """

        completion = client.chat.completions.create(
            model="meta-llama/Meta-Llama-3-8B-Instruct",
            messages=[
                {"role": "system", "content": "You are a resume expert."},
                {"role": "user", "content": prompt}
            ],
            max_tokens=250,
            temperature=0.7
        )

        result = completion.choices[0].message.content.strip()

        return JsonResponse({"result": result})

    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)


# skills
def generate_skills(request):

    if request.method != "POST":
        return JsonResponse({"error": "Only POST allowed"}, status=405)

    try:
        data = json.loads(request.body.decode("utf-8"))
        user_text = data.get("content", "").strip()

        if not user_text:
            return JsonResponse({"error": "Skills are empty"}, status=400)

        prompt = f"""
        Organize and improve these skills professionally:

        {user_text}

        Return only improved version.
        """

        completion = client.chat.completions.create(
            model="meta-llama/Meta-Llama-3-8B-Instruct",
            messages=[
                {"role": "system", "content": "You are a resume expert."},
                {"role": "user", "content": prompt}
            ],
            max_tokens=200,
            temperature=0.6
        )

        result = completion.choices[0].message.content.strip()

        return JsonResponse({"result": result})

    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)

# project

def generate_project(request):

    if request.method != "POST":
        return JsonResponse({"error": "Only POST allowed"}, status=405)

    try:
        data = json.loads(request.body.decode("utf-8"))
        user_text = data.get("content", "").strip()

        if not user_text:
            return JsonResponse({"error": "Project is empty"}, status=400)

        prompt = f"""
        Improve this project description professionally:

        {user_text}

        Return ONLY improved version.
        """

        completion = client.chat.completions.create(
            model="meta-llama/Meta-Llama-3-8B-Instruct",
            messages=[
                {"role": "system", "content": "You are a professional resume writer."},
                {"role": "user", "content": prompt}
            ],
            max_tokens=250,
            temperature=0.7
        )

        result = completion.choices[0].message.content.strip()

        return JsonResponse({"result": result})

    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)


# experience
def generate_work(request):

    if request.method != "POST":
        return JsonResponse({"error": "Only POST allowed"}, status=405)

    try:
        data = json.loads(request.body.decode("utf-8"))
        user_text = data.get("content", "").strip()

        if not user_text:
            return JsonResponse({"error": "Experience is empty"}, status=400)

        prompt = f"""
        Rewrite this job experience professionally with strong impact:

        {user_text}

        Return ONLY improved version.
        """

        completion = client.chat.completions.create(
            model="meta-llama/Meta-Llama-3-8B-Instruct",
            messages=[
                {"role": "system", "content": "You are a resume expert."},
                {"role": "user", "content": prompt}
            ],
            max_tokens=250,
            temperature=0.7
        )

        result = completion.choices[0].message.content.strip()

        return JsonResponse({"result": result})

    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)


from django.shortcuts import render
from users1.models import QuizQuestion, Topic


def quiz_home(request):

    topics = Topic.objects.all()

    return render(request, "quiz_home.html", {"topics": topics})


def start_quiz(request, topic_name):

    questions = QuizQuestion.objects.filter(subtopic__topic__name=topic_name)[:20]

    return render(request, "quiz.html", {
        "questions": questions,
        "topic": topic_name
    })


def result(request):

    return render(request, "result.html")